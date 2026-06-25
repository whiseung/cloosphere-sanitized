"""Admin Document Templates — admin-uploaded master templates for document tools.

PPT/Word/Excel 각 1개씩 글로벌 슬롯. `extension_modules/tools/document/{pptx,docx,xlsx}_tool.py`
가 빌더 진입부에서 이 config 를 읽어 마스터 슬라이드/스타일을 inherit 한다.

설계 결정 (dev/00_dev_docs/active/admin-document-templates/):
- 글로벌 (조직당 1세트), admin only — D1
- 확장자당 1개 슬롯 — D2
- 신규 라우터 `/api/v1/document-templates` — D5
- Storage provider 경유 (운영 Azure Blob auto-fallback) — D6
- License: `FeatureModule.DOCUMENT_TEMPLATES` (PROFESSIONAL+) — D7
"""

import asyncio
import io
import logging
import re
import time
import uuid
from typing import Optional

import httpx
from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
    status,
)
from fastapi.responses import StreamingResponse
from open_webui.constants import ERROR_MESSAGES
from open_webui.env import SRC_LOG_LEVELS
from open_webui.storage.provider import Storage
from open_webui.utils.audit_logger import AuditLogger
from open_webui.utils.auth import get_admin_user
from open_webui.utils.license import require_feature
from pydantic import BaseModel

log = logging.getLogger(__name__)
log.setLevel(SRC_LOG_LEVELS["MAIN"])

router = APIRouter()


# kind → (config attribute, MIME type, file extension)
KIND_MAP: dict[str, tuple[str, str, str]] = {
    "pptx": (
        "DOCUMENT_TEMPLATE_PPTX",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "pptx",
    ),
    "docx": (
        "DOCUMENT_TEMPLATE_DOCX",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "xlsx": (
        "DOCUMENT_TEMPLATE_XLSX",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
}

# OOXML zip magic — all .pptx/.docx/.xlsx start with PK\x03\x04
ZIP_MAGIC = b"PK\x03\x04"

# Schema version of the stored config dict (M1 — future migration guard)
CONFIG_SCHEMA_VERSION = 1


def _get_template_config(request: Request, kind: str) -> dict:
    """Return the stored dict for a kind (empty dict if none)."""
    config_attr = KIND_MAP[kind][0]
    val = getattr(request.app.state.config, config_attr, {})
    return val if isinstance(val, dict) else {}


def _public_meta(cfg: dict) -> dict:
    """Strip storage path before exposing to clients."""
    if not cfg:
        return {}
    return {
        "is_custom": True,
        "original_filename": cfg.get("original_filename", ""),
        "uploaded_at": cfg.get("uploaded_at", 0),
        "uploaded_by": cfg.get("uploaded_by", ""),
    }


def _trial_open(kind: str, raw: bytes) -> bool:
    """Open the bytes with the matching python-pptx/docx/openpyxl library.

    Returns True on success (template is a parseable file). False otherwise.
    Run BEFORE persisting the upload — fail fast so admin sees a clear error
    instead of silent runtime fallback later.
    """
    try:
        buf = io.BytesIO(raw)
        if kind == "pptx":
            from pptx import Presentation

            Presentation(buf)
        elif kind == "docx":
            from docx import Document

            Document(buf)
        elif kind == "xlsx":
            from openpyxl import load_workbook

            load_workbook(buf)
        else:
            return False
        return True
    except Exception as e:  # noqa: BLE001 — library raises many distinct types
        log.warning("Document template trial open failed for %s: %s", kind, e)
        return False


@router.get("/config", dependencies=[Depends(require_feature("document_templates"))])
async def get_document_templates_config(request: Request, user=Depends(get_admin_user)):
    """Return public metadata for all 3 slots — does NOT include file_path."""
    return {
        kind: _public_meta(_get_template_config(request, kind)) for kind in KIND_MAP
    }


@router.post(
    "/upload/{kind}", dependencies=[Depends(require_feature("document_templates"))]
)
async def upload_document_template(
    request: Request,
    kind: str,
    file: UploadFile,
    user=Depends(get_admin_user),
):
    if kind not in KIND_MAP:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(
                f"Invalid kind: {kind}. Expected one of {list(KIND_MAP)}"
            ),
        )

    config_attr, expected_mime, ext = KIND_MAP[kind]

    contents = await file.read()
    if not contents:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT("Empty file"),
        )

    # Magic-number check — defense against renamed non-Office files (L1)
    if not contents.startswith(ZIP_MAGIC):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(
                "File is not a valid Office Open XML (.pptx/.docx/.xlsx)"
            ),
        )

    # Extension check (defense in depth)
    if file.filename and not file.filename.lower().endswith(f".{ext}"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(f"Filename must end with .{ext}"),
        )

    # MIME header check (best-effort — content_type can be None)
    if file.content_type and file.content_type != expected_mime:
        log.info(
            "Document template upload MIME mismatch for %s: got %s, expected %s",
            kind,
            file.content_type,
            expected_mime,
        )

    # Trial open — corrupt templates rejected at upload time (M2)
    if not _trial_open(kind, contents):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(
                "Template file is corrupt or not a valid Office document"
            ),
        )

    # Stash old path for best-effort cleanup (HS2 — keep history minimal)
    old_cfg = _get_template_config(request, kind)
    old_path = old_cfg.get("file_path", "") if old_cfg else ""

    # Storage upload — unique key per upload so Azure ETag/lease race is avoided
    storage_key = f"document-templates/{kind}/{uuid.uuid4()}.{ext}"
    buffer = io.BytesIO(contents)
    _, file_path = Storage.upload_file(buffer, storage_key)

    # Persist config (AppConfig.__setattr__ → DB save + Redis pub/sub)
    new_value = {
        "version": CONFIG_SCHEMA_VERSION,
        "file_path": file_path,
        "original_filename": file.filename or f"template.{ext}",
        "uploaded_at": int(time.time()),
        "uploaded_by": user.id,
    }
    setattr(request.app.state.config, config_attr, new_value)

    # Best-effort cleanup of previous version (after new config is committed)
    if old_path and old_path != file_path:
        try:
            Storage.delete_file(old_path)
        except Exception as e:  # noqa: BLE001 — non-critical cleanup
            log.warning("Failed to clean previous template %s: %s", old_path, e)

    AuditLogger.log_settings_change(
        f"document-templates/upload/{kind}",
        after_data={
            "kind": kind,
            "original_filename": new_value["original_filename"],
        },
    )
    return _public_meta(new_value)


@router.get(
    "/{kind}/download",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def download_document_template(
    request: Request, kind: str, user=Depends(get_admin_user)
):
    """Admin verification — stream the stored template back. Admin-only."""
    if kind not in KIND_MAP:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(f"Invalid kind: {kind}"),
        )

    cfg = _get_template_config(request, kind)
    file_path: Optional[str] = cfg.get("file_path") if cfg else None
    if not file_path:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=ERROR_MESSAGES.NOT_FOUND,
        )

    _, mime, ext = KIND_MAP[kind]
    # Storage.get_file 은 모든 프로바이더에서 로컬 경로(str)를 반환 — bytes 아님.
    try:
        local_path = Storage.get_file(file_path)
        with open(local_path, "rb") as f:
            raw = f.read()
    except Exception as e:  # noqa: BLE001 — storage may raise provider-specific
        log.error("Failed to read template %s from storage: %s", file_path, e)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=ERROR_MESSAGES.DEFAULT("Failed to read template from storage"),
        ) from e

    original = cfg.get("original_filename") or f"template.{ext}"
    return StreamingResponse(
        io.BytesIO(raw),
        media_type=mime,
        headers={
            "Content-Disposition": f'attachment; filename="{original}"',
            "Cache-Control": "no-cache, must-revalidate",
        },
    )


@router.delete("/{kind}", dependencies=[Depends(require_feature("document_templates"))])
async def delete_document_template(
    request: Request, kind: str, user=Depends(get_admin_user)
):
    if kind not in KIND_MAP:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(f"Invalid kind: {kind}"),
        )

    config_attr = KIND_MAP[kind][0]
    cfg = _get_template_config(request, kind)
    file_path = cfg.get("file_path", "") if cfg else ""

    # Best-effort storage cleanup
    if file_path:
        try:
            Storage.delete_file(file_path)
        except Exception as e:  # noqa: BLE001
            log.warning("Failed to delete template file %s: %s", file_path, e)

    # Reset config — empty dict means "no template configured"
    setattr(request.app.state.config, config_attr, {})

    AuditLogger.log_settings_change(
        f"document-templates/delete/{kind}",
        after_data={"kind": kind, "cleared": True},
    )
    return {"status": "ok"}


# ────────────────────────────────────────────────────────────────────────────
# Presenton — PPT 생성 엔진 연결 설정 (관리자 UI: 문서 템플릿 탭)
# enabled 시 document_tools 의 PPT 경로가 내장 python-pptx → Presenton 으로 대체된다.
# config 는 PRESENTON_* PersistentConfig (main.py 에서 app.state.config 등록).
# ────────────────────────────────────────────────────────────────────────────


# 연결(base_url/timeout)은 마켓플레이스(관리자>설정>마켓플레이스, PPT Generator)가
# 관리한다. 이 탭은 "엔진 사용 여부(enabled)"와 "기본 템플릿(default_template)"만 결정.
# (기존 클라이언트가 base_url/timeout 을 보내도 Pydantic extra-ignore 로 무시됨.)
class PresentonConfigForm(BaseModel):
    enabled: bool
    default_template: str = "general"


def _presenton_base(request: Request, override: Optional[str] = None) -> str:
    base = (
        override
        or request.app.state.config.PRESENTON_BASE_URL
        or "http://localhost:5001"
    )
    return base.rstrip("/")


@router.get(
    "/presenton/config",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def get_presenton_config(request: Request, user=Depends(get_admin_user)):
    c = request.app.state.config
    return {
        "enabled": bool(c.PRESENTON_ENABLED),
        "base_url": c.PRESENTON_BASE_URL or "",
        "timeout": int(c.PRESENTON_TIMEOUT or 600),
        "default_template": c.PRESENTON_DEFAULT_TEMPLATE or "general",
    }


@router.post(
    "/presenton/config",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def update_presenton_config(
    request: Request, form_data: PresentonConfigForm, user=Depends(get_admin_user)
):
    c = request.app.state.config
    c.PRESENTON_ENABLED = bool(form_data.enabled)
    c.PRESENTON_DEFAULT_TEMPLATE = form_data.default_template or "general"
    AuditLogger.log_settings_change(
        "document-templates/presenton", after_data=form_data.model_dump()
    )
    return await get_presenton_config(request, user)


@router.get(
    "/presenton/test",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def test_presenton_connection(
    request: Request, base_url: Optional[str] = None, user=Depends(get_admin_user)
):
    """연결 테스트 — Presenton /template/all 조회. base_url 쿼리로 미저장 값 테스트 가능."""
    base = _presenton_base(request, base_url)
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            r = await client.get(f"{base}/api/v1/ppt/template/all")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(f"Presenton 연결 실패 ({base}): {e}"),
        )
    if r.status_code != 200:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT(
                f"Presenton 응답 오류 (HTTP {r.status_code})"
            ),
        )
    items = r.json()
    count = len(items) if isinstance(items, list) else 0
    return {"ok": True, "template_count": count}


@router.get(
    "/presenton/templates",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def list_presenton_templates(request: Request, user=Depends(get_admin_user)):
    """사용 가능 템플릿(내장+커스텀) 목록 — 기본 템플릿 드롭다운용. 실패 시 빈 목록."""
    base = _presenton_base(request)
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            r = await client.get(f"{base}/api/v1/ppt/template/all")
        items = r.json() if r.status_code == 200 else []
        return [
            {"id": str(t.get("id")), "name": str(t.get("name") or t.get("id"))}
            for t in items
            if isinstance(t, dict) and t.get("id")
        ]
    except Exception as e:  # noqa: BLE001
        log.warning("Presenton 템플릿 목록 조회 실패: %s", e)
        return []


# ── PPT → Presenton 커스텀 템플릿 추출 (업로드 또는 기존 마스터 PPT 재사용) ──────
# 추출은 슬라이드당 vision-LLM ~60-90s 라 백그라운드 잡 + 폴링. (단일 프로세스 dev 기준
# in-memory 레지스트리 — 멀티워커 배포 시 공유 스토어 필요.)
_PRESENTON_JOBS: dict[str, dict] = {}
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _layout_meta(rc: str, idx: int) -> tuple[str, str]:
    lid = re.search(r'const layoutId\s*=\s*"([^"]+)"', rc)
    lname = re.search(r'const layoutName\s*=\s*"([^"]+)"', rc)
    return (lid.group(1) if lid else f"layout-{idx}"), (
        lname.group(1) if lname else f"Layout {idx}"
    )


async def _build_presenton_template(
    base: str,
    pptx_bytes: bytes,
    filename: str,
    name: str,
    indices: list[int],
    job_id: str,
) -> None:
    """업로드/프리뷰 → init → 슬라이드별 레이아웃 추출 → save. 진행상황을 잡에 기록."""
    job = _PRESENTON_JOBS[job_id]
    try:
        async with httpx.AsyncClient(timeout=300.0) as client:
            job["message"] = "슬라이드 프리뷰 생성 중…"
            files = {"pptx_file": (filename, pptx_bytes, _PPTX_MIME)}
            r = await client.post(
                f"{base}/api/v1/ppt/template/fonts-upload-and-slides-preview",
                files=files,
                timeout=120.0,
            )
            r.raise_for_status()
            prev = r.json()
            n = len(prev.get("slide_image_urls", []))
            idx = [i for i in indices if 0 <= i < n] or [0]

            job["message"] = "템플릿 초기화 중… (PPTX→HTML)"
            r = await client.post(
                f"{base}/api/v1/ppt/template/create/init",
                json={
                    "slide_image_urls": prev["slide_image_urls"],
                    "pptx_url": prev["pptx_url"],
                    "fonts": prev.get("fonts"),
                },
                timeout=180.0,
            )
            r.raise_for_status()
            tcid = str(r.json()).strip('"')

            layouts = []
            for k, i in enumerate(idx):
                job["message"] = f"레이아웃 추출 중… ({k + 1}/{len(idx)})"
                r = await client.post(
                    f"{base}/api/v1/ppt/template/slide-layout/create",
                    json={"id": tcid, "index": i},
                    timeout=300.0,
                )
                r.raise_for_status()
                rc = r.json().get("react_component", "")
                lid, lname = _layout_meta(rc, i)
                layouts.append(
                    {"layout_id": lid, "layout_name": lname, "layout_code": rc}
                )

            job["message"] = "템플릿 저장 중…"
            r = await client.post(
                f"{base}/api/v1/ppt/template/save",
                json={
                    "template_info_id": tcid,
                    "name": name,
                    "description": f"Extracted from {filename}",
                    "layouts": layouts,
                },
                timeout=60.0,
            )
            r.raise_for_status()
            tid = r.json().get("id")
            job.update(
                {"status": "completed", "message": "완료", "template_id": str(tid)}
            )
    except Exception as e:  # noqa: BLE001 — 어떤 단계 실패든 잡에 기록
        log.error("Presenton 템플릿 추출 실패 (job=%s): %s", job_id, e)
        job.update({"status": "failed", "message": "실패", "error": str(e)[:300]})


@router.post(
    "/presenton/templates/create",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def create_presenton_template(
    request: Request,
    name: str = Form(...),
    indices: str = Form("0,1,2,3,4"),
    from_master: bool = Form(False),
    file: Optional[UploadFile] = File(None),
    user=Depends(get_admin_user),
):
    """회사 PPT(업로드 또는 기존 마스터)를 Presenton 커스텀 템플릿으로 추출 (백그라운드 잡)."""
    if from_master:
        cfg = _get_template_config(request, "pptx")
        fp = cfg.get("file_path") if cfg else None
        if not fp:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=ERROR_MESSAGES.DEFAULT("업로드된 PPT 마스터 템플릿이 없습니다."),
            )
        try:
            local = Storage.get_file(fp)
            with open(local, "rb") as f:
                pptx_bytes = f.read()
        except Exception as e:  # noqa: BLE001
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=ERROR_MESSAGES.DEFAULT(f"마스터 PPT 읽기 실패: {e}"),
            ) from e
        filename = cfg.get("original_filename") or "master.pptx"
    elif file is not None:
        pptx_bytes = await file.read()
        filename = file.filename or "upload.pptx"
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=ERROR_MESSAGES.DEFAULT("pptx 파일 또는 from_master 가 필요합니다."),
        )

    try:
        idx = [int(x) for x in indices.split(",") if x.strip()]
    except ValueError:
        idx = [0, 1, 2]

    base = _presenton_base(request)
    job_id = str(uuid.uuid4())
    _PRESENTON_JOBS[job_id] = {
        "status": "pending",
        "message": "대기 중…",
        "template_id": None,
        "error": None,
    }
    asyncio.create_task(
        _build_presenton_template(base, pptx_bytes, filename, name, idx, job_id)
    )
    AuditLogger.log_settings_change(
        "document-templates/presenton/create",
        after_data={"name": name, "indices": idx, "from_master": from_master},
    )
    return {"job_id": job_id}


@router.get(
    "/presenton/templates/job/{job_id}",
    dependencies=[Depends(require_feature("document_templates"))],
)
async def get_presenton_template_job(
    request: Request, job_id: str, user=Depends(get_admin_user)
):
    job = _PRESENTON_JOBS.get(job_id)
    if job is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail=ERROR_MESSAGES.NOT_FOUND
        )
    return job
