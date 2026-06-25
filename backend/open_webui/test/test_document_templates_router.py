"""Document Templates router unit tests — helper functions.

Full-flow tests (upload→download round-trip with auth + require_feature)
would require fixturing the entire FastAPI app + DB + license enforcement
state, which is out of proportion to the routing logic being tested. The
critical pieces (trial_open, magic-number check, public_meta strip) are
covered as pure unit tests below; end-to-end flows are validated by the
builder-side integration tests in
``extension_modules/tools/document/tests/test_template_loading.py``.
"""

from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


@pytest.fixture
def valid_pptx_bytes() -> bytes:
    prs = Presentation()
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def valid_docx_bytes() -> bytes:
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def valid_xlsx_bytes() -> bytes:
    wb = Workbook()
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestTrialOpen:
    """_trial_open should accept real Office files and reject anything else."""

    def test_accepts_real_pptx(self, valid_pptx_bytes):
        from open_webui.routers.document_templates import _trial_open

        assert _trial_open("pptx", valid_pptx_bytes) is True

    def test_accepts_real_docx(self, valid_docx_bytes):
        from open_webui.routers.document_templates import _trial_open

        assert _trial_open("docx", valid_docx_bytes) is True

    def test_accepts_real_xlsx(self, valid_xlsx_bytes):
        from open_webui.routers.document_templates import _trial_open

        assert _trial_open("xlsx", valid_xlsx_bytes) is True

    def test_rejects_random_bytes(self):
        from open_webui.routers.document_templates import _trial_open

        assert _trial_open("pptx", b"not a zip file") is False

    def test_rejects_truncated_zip(self):
        from open_webui.routers.document_templates import _trial_open

        # ZIP magic but body cut off
        assert _trial_open("pptx", b"PK\x03\x04" + b"\x00" * 50) is False

    def test_rejects_kind_xlsx_with_pptx_bytes(self, valid_pptx_bytes):
        """pptx is also a valid zip — but xlsx loader rejects it."""
        from open_webui.routers.document_templates import _trial_open

        # openpyxl accepts the zip but raises on missing xl/workbook.xml
        result = _trial_open("xlsx", valid_pptx_bytes)
        # openpyxl may raise InvalidFileException or similar — _trial_open
        # catches and returns False
        assert result is False

    def test_unknown_kind_returns_false(self, valid_pptx_bytes):
        from open_webui.routers.document_templates import _trial_open

        assert _trial_open("unknown", valid_pptx_bytes) is False


class TestPublicMeta:
    def test_empty_dict_returns_empty(self):
        from open_webui.routers.document_templates import _public_meta

        assert _public_meta({}) == {}

    def test_strips_file_path(self):
        from open_webui.routers.document_templates import _public_meta

        cfg = {
            "version": 1,
            "file_path": "document-templates/pptx/abc.pptx",
            "original_filename": "월간보고서.pptx",
            "uploaded_at": 1700000000,
            "uploaded_by": "u-1",
        }
        meta = _public_meta(cfg)
        assert "file_path" not in meta
        assert meta["original_filename"] == "월간보고서.pptx"
        assert meta["uploaded_at"] == 1700000000
        assert meta["uploaded_by"] == "u-1"
        assert meta["is_custom"] is True

    def test_handles_missing_fields_gracefully(self):
        from open_webui.routers.document_templates import _public_meta

        meta = _public_meta({"file_path": "x"})  # only file_path set
        assert meta["original_filename"] == ""
        assert meta["uploaded_at"] == 0
        assert meta["uploaded_by"] == ""


class TestKindMap:
    def test_three_kinds_present(self):
        from open_webui.routers.document_templates import KIND_MAP

        assert set(KIND_MAP.keys()) == {"pptx", "docx", "xlsx"}

    def test_mime_types_are_office_open_xml(self):
        from open_webui.routers.document_templates import KIND_MAP

        for kind, (config_attr, mime, ext) in KIND_MAP.items():
            assert config_attr == f"DOCUMENT_TEMPLATE_{kind.upper()}"
            assert "openxmlformats" in mime
            assert ext == kind


class TestZipMagic:
    def test_constant_value(self):
        from open_webui.routers.document_templates import ZIP_MAGIC

        # ZIP local file header — all OOXML files start with this
        assert ZIP_MAGIC == b"PK\x03\x04"

    def test_real_pptx_starts_with_magic(self, valid_pptx_bytes):
        from open_webui.routers.document_templates import ZIP_MAGIC

        assert valid_pptx_bytes.startswith(ZIP_MAGIC)

    def test_real_docx_starts_with_magic(self, valid_docx_bytes):
        from open_webui.routers.document_templates import ZIP_MAGIC

        assert valid_docx_bytes.startswith(ZIP_MAGIC)

    def test_real_xlsx_starts_with_magic(self, valid_xlsx_bytes):
        from open_webui.routers.document_templates import ZIP_MAGIC

        assert valid_xlsx_bytes.startswith(ZIP_MAGIC)
